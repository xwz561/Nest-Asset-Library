from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\Users\EDY\Documents\New project\Nest-Asset-Library-1.6.0-Source-xwz\folder-installer-3.0.2-beta.54-welcome-brand-icon\小旺仔素材库测试版使用教程 beta54.docx")


def set_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_paragraph(p, before=0, after=6, line=1.35):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_body(doc, text, before=0, after=7):
    p = doc.add_paragraph()
    style_paragraph(p, before, after)
    set_font(p.add_run(text), size=10.8, color=(35, 35, 35))
    return p


def add_step(doc, number, title, text):
    p = doc.add_paragraph()
    style_paragraph(p, before=3, after=3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"{number}. {title}")
    set_font(run, size=11, bold=True, color=(0, 0, 0))
    p2 = doc.add_paragraph()
    style_paragraph(p2, before=0, after=7)
    set_font(p2.add_run(text), size=10.6, color=(35, 35, 35))


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(13 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, size=15 if level == 1 else 12, bold=True, color=(0, 0, 0))
    return p


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.8)

    title_style = doc.styles["Title"]
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(title, before=0, after=4, line=1.0)
    set_font(title.add_run("小旺仔素材库测试版使用教程"), size=23, bold=True, color=(0, 0, 0))
    subtitle = doc.add_paragraph()
    style_paragraph(subtitle, before=0, after=14, line=1.0)
    set_font(subtitle.add_run("适用于 3 0 2 beta 54"), size=11, color=(80, 80, 80))

    intro = doc.add_paragraph()
    style_paragraph(intro, before=0, after=10)
    lead = intro.add_run("本教程用于安装和验证小旺仔素材库 beta.54。")
    set_font(lead, size=10.8, bold=True, color=(0, 0, 0))
    set_font(intro.add_run(" 该版本会在侧栏显示准确的版本号和“测试版”标记，并包含 AI Flow 素材补关联与安全删除同步功能。"), size=10.8, color=(35, 35, 35))

    add_heading(doc, "开始前确认", 1)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.72)
    table.columns[1].width = Inches(4.9)
    head = table.rows[0].cells
    for cell, text in zip(head, ("项目", "确认内容")):
        shade(cell, "1F4E79")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), size=10.4, bold=True, color=(255, 255, 255))
    for left, right in [
        ("安装包", "使用小旺仔素材库安装包 3.0.2 beta.54.exe。该包仅供本机测试，未做代码签名。"),
        ("网页浏览器", "使用 Microsoft Edge，保持 AI Flow 页面已登录。默认服务地址为 http://10.128.20.135:8080/。"),
        ("历史包", "安装 beta.54 不会清理 beta.53、便携版或原有素材库。不要手动删除正在使用的素材库目录。"),
    ]:
        row = table.add_row().cells
        for index, text in enumerate((left, right)):
            set_cell_border(row[index])
            set_cell_margins(row[index])
            row[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = row[index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(text), size=10.2, bold=(index == 0), color=(30, 30, 30))

    add_heading(doc, "安装测试版", 1)
    add_step(doc, 1, "关闭旧版本", "先退出正在运行的小旺仔素材库。若任务栏仍有程序图标，请确认窗口和后台进程都已关闭。")
    add_step(doc, 2, "运行安装包", "双击安装包并等待完成。Windows 可能显示未知发布者提示，这是未签名的本地测试包；确认文件路径正确后，再选择继续运行。")
    add_step(doc, 3, "打开软件", "安装完成后启动小旺仔素材库。现有资源库数据会继续使用，不需要重新导入素材。")

    add_heading(doc, "确认版本显示", 1)
    add_body(doc, "在软件左下角磁盘容量下方，应该显示“v3.0.2-beta.54”和橙色“测试版”标记。若仍显示 beta.14，说明打开的仍是旧程序，请完全退出旧版本后从新安装的快捷方式启动。")

    add_heading(doc, "安装并启用 Edge 扩展", 1)
    add_step(doc, 1, "打开扩展设置", "在软件顶部的“扩展”中选择 Microsoft Edge 的“准备”。软件会打开本地扩展目录，并尝试打开 Edge 的扩展管理页面。")
    add_step(doc, 2, "加载扩展", "在 Edge 的 edge://extensions 页面开启“开发人员模式”，选择“加载解压缩的扩展”，然后选择软件刚打开的本地扩展目录。")
    add_step(doc, 3, "刷新 AI Flow", "保持 AI Flow 页面已登录并按一次刷新。扩展重新载入后，网页与素材库才能接收实时同步、上传和删除请求。")

    add_heading(doc, "开启实时同步", 1)
    add_body(doc, "在需要同步的本地文件夹上右键，选择上传到 AI Flow 或开启 AI Flow 实时同步。已开启的同步根目录及其子文件夹会显示闪电图标。闪电图标仅代表该目录已启用同步，不表示每一张历史素材都已经建立服务器关联。")
    add_body(doc, "本地新增素材后，保留 Edge 的 AI Flow 页面打开。素材卡片可能短暂显示“等待上传”；上传完成后，网页端可刷新查看。较大的文件会受网络和网页服务处理时间影响。")

    add_heading(doc, "补关联旧素材", 1)
    add_body(doc, "测试文件夹中早期导入的素材，可能没有蓝色“来自 AI Flow”标记。这类素材过去没有保存服务器素材 ID，即使服务器已有同内容文件，也不能直接执行服务器删除。")
    add_step(doc, 1, "保持网页已登录", "打开并刷新目标 AI Flow 项目的网页，保持 Edge 窗口不关闭。")
    add_step(doc, 2, "等待同步或手动同步", "等待实时同步回拉，或在网页素材区域使用“同步目录和素材”。beta.54 会对同一同步文件夹中的文件做内容哈希匹配。")
    add_step(doc, 3, "确认结果", "内容相同的旧本地素材不会复制第二份，也不会重新上传；它会补写服务器素材 ID，并显示蓝色“来自 AI Flow”标记。")

    add_heading(doc, "本地删除同步到服务器", 1)
    add_body(doc, "只删除已经显示蓝色“来自 AI Flow”标记的素材。将它从素材库删除后，软件会先移入本地回收站，并让已登录的 Edge 扩展请求删除同一服务器素材。服务器确认成功后，网页刷新即可看到素材消失。")
    add_body(doc, "没有蓝色标记的历史素材不会盲删服务器文件。这是为了避免同名素材或其他项目素材被误删。请先按上一节完成补关联，再测试删除。")

    add_heading(doc, "常见问题", 1)
    faq = [
        ("显示等待上传很久", "确认 Edge 已登录 AI Flow、扩展处于启用状态，并刷新 AI Flow 页面。大文件还需要等待网页服务接收和生成缩略图。"),
        ("网页没有新素材", "先确认本地素材所在目录已开启实时同步。然后保持 AI Flow 页面打开并刷新；如仍无变化，检查扩展是否因更新而需要重新加载。"),
        ("本地删了网页还在", "检查被删卡片在删除前是否有蓝色“来自 AI Flow”标记。没有服务器 ID 的旧素材不会发送删除请求；先完成补关联。"),
    ]
    for question, answer in faq:
        p = doc.add_paragraph()
        style_paragraph(p, before=3, after=2)
        set_font(p.add_run(question), size=10.8, bold=True, color=(0, 0, 0))
        p2 = doc.add_paragraph()
        style_paragraph(p2, before=0, after=6)
        set_font(p2.add_run(answer), size=10.5, color=(35, 35, 35))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("小旺仔素材库 beta.54 测试版"), size=8.5, color=(100, 100, 100))
    doc.core_properties.title = "小旺仔素材库测试版使用教程"
    doc.core_properties.subject = "小旺仔素材库 beta.54 安装和 AI Flow 同步教程"
    doc.core_properties.author = "小旺仔素材库"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
